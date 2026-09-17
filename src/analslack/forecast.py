"""세일즈 포케스트 엑셀 업로드 파싱.

업로드된 엑셀의 컬럼 구조를 미리 고정하지 않는다 — 첫 번째 행을 그대로
헤더로 쓰고, 그 아래 행들을 데이터로 읽는다. 워크북의 첫 번째(활성) 시트만
사용한다.
"""

from __future__ import annotations

import datetime
from typing import Any, BinaryIO

from openpyxl import load_workbook

MAX_ROWS = 20000


class ForecastParseError(ValueError):
    pass


def _jsonify_cell(value: Any) -> Any:
    """openpyxl 셀 값을 JSON으로 저장/렌더링 가능한 값으로 변환."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime.datetime):
        # 순수 날짜 셀도 openpyxl은 datetime으로 돌려주므로, 시각이 자정이면
        # 날짜만(YYYY-MM-DD) 남긴다 — 대부분의 포케스트 날짜 컬럼(마감일 등)에
        # 시각까지 있는 경우는 드물다.
        if value.time() == datetime.time(0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return str(value)


def parse_excel(file_stream: BinaryIO) -> "tuple[list[str], list[list[Any]]]":
    """엑셀 파일(바이너리 스트림)을 (컬럼 목록, 행 목록)으로 파싱한다.

    첫 행을 헤더로 사용하며, 완전히 빈 행은 건너뛴다.
    """
    try:
        workbook = load_workbook(file_stream, read_only=True, data_only=True)
    except Exception as e:
        raise ForecastParseError(f"엑셀 파일을 읽을 수 없습니다: {e}") from e

    try:
        sheet = workbook.active
        if sheet is None:
            raise ForecastParseError("엑셀 파일에 시트가 없습니다.")

        row_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(row_iter)
        except StopIteration:
            raise ForecastParseError("엑셀 파일이 비어 있습니다.")

        columns = [
            (str(c).strip() if c is not None and str(c).strip() else f"열{i + 1}")
            for i, c in enumerate(header_row)
        ]
        if not columns:
            raise ForecastParseError("헤더(첫 행)에 컬럼이 없습니다.")

        rows: list[list[Any]] = []
        for raw_row in row_iter:
            if raw_row is None or all(c is None for c in raw_row):
                continue
            row = [_jsonify_cell(c) for c in raw_row[: len(columns)]]
            if len(row) < len(columns):
                row.extend([None] * (len(columns) - len(row)))
            rows.append(row)
            if len(rows) >= MAX_ROWS:
                break

        return columns, rows
    finally:
        workbook.close()
